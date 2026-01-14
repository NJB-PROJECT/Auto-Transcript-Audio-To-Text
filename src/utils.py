import os
from datetime import datetime
from docx import Document

def save_to_txt(controls, filename=None):
    """
    Menyimpan isi transcript (list of Flet controls) ke file TXT.
    """
    if not filename:
        filename = f"transcript_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"

    content = ""
    for control in controls:
        # Assuming control is Container -> Text or just Text
        if hasattr(control, "content") and hasattr(control.content, "value"):
            content += control.content.value + "\n\n"
        elif hasattr(control, "value"):
            content += control.value + "\n\n"

    with open(filename, "w", encoding="utf-8") as f:
        f.write(content)
    return filename

def save_to_docx(controls, filename=None):
    """
    Menyimpan isi transcript ke file Word (.docx).
    """
    if not filename:
        filename = f"transcript_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"

    doc = Document()
    doc.add_heading('Transcript Audio', 0)

    for control in controls:
        text = ""
        if hasattr(control, "content") and hasattr(control.content, "value"):
            text = control.content.value
        elif hasattr(control, "value"):
            text = control.value

        if text:
            # Clean up UI markers if any
            clean_text = text.replace("• ", "")
            p = doc.add_paragraph(clean_text)

    doc.save(filename)
    return filename
